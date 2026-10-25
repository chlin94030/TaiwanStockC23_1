"""V23.1 integrated pipeline: pinned V22 market cache -> dated PIT factors ->
fixed research hypotheses -> nonoverlapping daily NAV -> benchmark report.

No automatic orders, never opens V22 final holdout, no network in validation.
A running worker may be interrupted by the Streamlit server lifecycle; immutable
input caches and completed job artifacts persist on the same host, and the user
may explicitly resume. Streamlit local disk is NOT persistent cloud storage.
"""
from __future__ import annotations
import hashlib
import json
import threading
import time
import traceback
import zipfile
from dataclasses import asdict
from datetime import datetime,timezone
from pathlib import Path
import numpy as np
import pandas as pd
from lab_io import DataError,atomic_json,read_json,sha
from studio_pool import digest,load_dataset,discover_prepared
from studio_cache import load_labels,store_for_dataset
from v22_cache import prepare_v22
from v22_runner import configure, DEFAULT_SPEC
from v231_data import evidence_signature,load_events
from v231_factors import build_asof_fundamentals,benchmark_rs120
from v231_engine import MODELS,V23Config,VERSION,select_asof,choose_cohorts,simulate_nav
from v231_release import load as load_release

_ACTIVE={};_ACTIVE_LOCK=threading.Lock()


def job_path(root):return Path(root)/'v23'/'latest.json'


def get_state(root):
    latest=read_json(job_path(root))
    if not latest.get('run'):return {}
    p=Path(latest['run'])
    if not p.is_relative_to(Path(root).resolve()):return {}
    return read_json(p/'state.json')


def v231_busy(root):
    with _ACTIVE_LOCK:
        worker=_ACTIVE.get(str(Path(root).resolve()))
        return bool(worker and worker.is_alive())


def _stamp(run,**changes):
    s=read_json(run/'state.json');s.update(changes,at=datetime.now(timezone.utc).isoformat())
    atomic_json(run/'state.json',s)


def _read_base(root,cfg,progress):
    folder,meta=load_dataset(root)
    progress('PINNED_DATA','檢查已固定的 V22 行情及母池（不重新下載）')
    feature_dir,info,cache_hit=prepare_v22(root,lambda **kw:progress('FEATURE_CACHE',str(kw.get('message','V22特徵建置中'))),lambda:None)
    if cfg.pool_rule not in info.get('pools',[]):raise DataError('指定的KOSTO250母池不在既有快取')
    vcfg,vspec=configure(meta,{**DEFAULT_SPEC,'pool_rule':cfg.pool_rule,'top_k':cfg.top_k,'capital':cfg.capital})
    labels,label_meta=load_labels(root,feature_dir,vcfg,lambda **kw:progress('LABEL_CACHE',str(kw.get('message','歷史報酬標籤快取'))),lambda:None)
    snap=pd.read_csv(feature_dir/'features.csv.gz',parse_dates=['date'])
    pools=pd.read_csv(feature_dir/'pools.csv.gz',parse_dates=['date'])
    stages=pd.read_csv(feature_dir/'market_stage.csv.gz',parse_dates=['date'])
    cal=pd.DatetimeIndex(read_json(folder/'calendar.json'))
    store=store_for_dataset(root,meta)
    bench=store.get('adj','0050.TW',meta['source_config']['start'],meta['source_config']['end'])
    if bench.empty:raise DataError('既有資料庫缺少0050的調整後價量資料，請先還原檢查點')
    bench=bench[['date','open','close']]
    return {'folder':folder,'meta':meta,'features':snap,'pools':pools,'stages':stages,'labels':labels,
            'cal':cal,'store':store,'bench':bench,'info':info,'labels_info':label_meta,
            'cache_hit':cache_hit}


def _research_dates(base,cfg):
    pool=base['pools'];dates=pd.DatetimeIndex(sorted(pool.loc[pool.pool==cfg.pool_rule,'date'].drop_duplicates()))
    if len(dates)<cfg.min_dev_train_months+cfg.final_holdout_months+8:
        raise DataError('樣本月份不足以同時分出訓練、開發驗證與未動的最後24個月')
    reserved=dates[-cfg.final_holdout_months:]
    # Never inspect any label on or after the reserved tail boundary; additionally
    # use each H-specific exit before boundary to prevent boundary leakage.
    dev=dates[cfg.min_dev_train_months:-cfg.final_holdout_months]
    if len(dev)<8:raise DataError('只有不足8個開發期月份，禁止挑選冠軍')
    return dates,dev,{'reserved_start':str(reserved[0].date()),'reserved_months':len(reserved),
                      'first_development_signal':str(dev[0].date()),'last_development_signal':str(dev[-1].date()),
                      'final_tail_inspected':False}


def _prepare_features(base,events,dev_dates,cfg,progress):
    snap=base['features'];pools=base['pools'];stages=base['stages'];dates=pd.DatetimeIndex(dev_dates)
    p=pools[(pools.pool==cfg.pool_rule)&(pools.date.isin(dates))][['date','ticker','amount20']]
    panel=snap[snap.date.isin(dates)].merge(p,on=['date','ticker'],how='inner',suffixes=('','_pool'),validate='one_to_one')
    if 'amount20_pool' in panel:
        panel['amount20']=panel['amount20_pool']
        panel=panel.drop(columns=['amount20_pool'])
    if panel.empty:raise DataError('開發期所選母池的歷史特徵不存在')
    panel=benchmark_rs120(panel,base['bench'],dates)
    context=stages[(stages.pool==cfg.pool_rule)&(stages.date.isin(dates))][['date','stage','stage_confidence']]
    panel=panel.merge(context,on='date',how='left',validate='many_to_one')
    chunks=[]
    for i,date in enumerate(dates):
        if i%5==0:progress('PIT_JOIN',f'依公告時間建立基本面訊號 {i+1}/{len(dates)}')
        sub=panel[panel.date==date].copy()
        f=build_asof_fundamentals(events,date)
        sub=sub.merge(f,on='ticker',how='left',validate='one_to_one')
        chunks.append(sub)
    panel=pd.concat(chunks,ignore_index=True)
    return panel


def _decision_table(base,panel,dev_dates,cfg,progress,tail_boundary,models):
    labels=base['labels']
    stages=panel[['date','stage','stage_confidence']].drop_duplicates('date')
    # label lookups on exposed dev only, never on the final reserved dates.
    labels=labels[(labels.date.isin(dev_dates)) & (labels.exit < pd.Timestamp(tail_boundary))].copy()
    b=labels[labels.ticker=='0050.TW'].set_index(['date','horizon'])
    stock=labels[labels.ticker!='0050.TW'].set_index(['date','horizon','ticker'])
    rows=[];blocked=0
    for ix,date in enumerate(dev_dates):
        group=panel[panel.date==date]
        if len(group)<cfg.top_k:continue
        stage=str(group.stage.iloc[0])
        for model in models:
            sel=select_asof(group,model,stage,cfg,models=models)
            for h in (5,20,60):
                if (date,h) not in b.index:continue
                benchmark=b.loc[(date,h)]
                if pd.isna(benchmark.get('exit')) or pd.isna(benchmark.get('entry')):
                    continue
                picks=list(sel.get('selected',[]));state=str(sel['state'])
                effective=float(sel['target_stock_fraction']);traded=[]
                if effective>0:
                    missing=False
                    for item in picks:
                        key=(date,h,item['ticker'])
                        if key not in stock.index:missing=True;break
                        lab=stock.loc[key]
                        if pd.isna(lab.net) or lab.state not in ('ASSUMED_OPEN_FILL','PREDECLARED_GAP_NO_FILL','NO_TRADE_VOLUME'):
                            missing=True;break
                        if bool(lab.get('filled',False)) and lab.state=='ASSUMED_OPEN_FILL':traded.append(item)
                    if missing:
                        state='EXECUTION_LABEL_BLOCKED';blocked+=1;effective=0.;traded=[]
                    else:effective*=len(traded)/max(len(picks),1)
                rows.append({'date':date,'horizon':h,'model':model,'entry':pd.Timestamp(benchmark.entry),
                             'exit':pd.Timestamp(benchmark.exit),'stage':stage,'eligible':sel.get('eligible',0),
                             'factor_coverage':sel.get('coverage',1.),'state':state,
                             'selected_count':len(picks),'filled_count':len(traded),
                             'target_stock_fraction':float(effective),'selected':traded,
                             'benchmark_label_net':float(benchmark.net) if pd.notna(benchmark.net) else np.nan})
        if ix%8==0:progress('SELECT',f'建置 5/20/60 日模擬持股 {ix+1}/{len(dev_dates)}')
    return pd.DataFrame(rows),blocked


def _subset_prices(base,decisions):
    wanted=set()
    for p in decisions.loc[decisions.target_stock_fraction>0,'selected']:
        wanted.update(o['ticker'] for o in p)
    store=base['store'];meta=base['meta'];cfg=meta['source_config']
    return {ticker:store.get('adj',ticker,cfg['start'],cfg['end'])[['date','open','close']] for ticker in wanted}


def _daily_ci(daily,seed=2312026,block=21,samples=400):
    if daily is None or len(daily)<100:return (None,None)
    x=np.log(daily.strategy_nav.to_numpy(float)[1:]/daily.strategy_nav.to_numpy(float)[:-1])-\
      np.log(daily.benchmark_nav.to_numpy(float)[1:]/daily.benchmark_nav.to_numpy(float)[:-1])
    n=len(x);rng=np.random.default_rng(seed);reps=[]
    for _ in range(samples):
        starts=rng.integers(0,n,size=int(np.ceil(n/block)))
        indices=np.concatenate([(np.arange(s,s+block)%n) for s in starts])[:n]
        reps.append(100*(np.exp(np.mean(x[indices])*252)-1))
    return tuple(float(v) for v in np.percentile(reps,[2.5,97.5]))


def _validate_one(base,decisions,prices,model,h,cfg,progress):
    data=decisions[(decisions.model==model)&(decisions.horizon==h)].sort_values('date')
    if data.empty:
        return {'status':'NOT_ENOUGH_MATURED_SIGNALS','model':model,'horizon':h,'decision_count':0},pd.DataFrame()
    # Restrict evaluation to cases with completely matured exits before reserved
    # boundary; filter occurs before any metrics or strategy selection.
    selected=choose_cohorts(data,h)
    if selected.empty:
        return {'status':'NO_NONOVERLAP_COHORTS','model':model,'horizon':h},pd.DataFrame()
    start=pd.Timestamp(selected.iloc[0]['date'])
    end=pd.Timestamp(selected.iloc[-1]['exit'])
    cal=base['cal'];cl=cal[(cal>=start)&(cal<=end)]
    if len(cl)<40:
        return {'status':'TOO_FEW_SESSIONS','model':model,'horizon':h},pd.DataFrame()
    try:
        daily,mt=simulate_nav(cl,base['bench'],prices,selected,cfg)
        _,mt_stress=simulate_nav(cl,base['bench'],prices,selected,cfg,stress=True)
        ci=_daily_ci(daily)
        n=int(selected.active_cohort.sum())
        blocked=int((data.state=='EXECUTION_LABEL_BLOCKED').sum())
        fund_block=int((data.state=='FUNDAMENTAL_COVERAGE_BLOCKED').sum())
        other_block=int((~data.state.isin(['PICKS_READY','ETF_BASELINE','EXECUTION_LABEL_BLOCKED','FUNDAMENTAL_COVERAGE_BLOCKED'])).sum())
        cover=1.-float(blocked+fund_block+other_block)/max(len(data),1) if model!='0050_ONLY' else 1.
        verdict='DEVELOPMENT_DIAGNOSTIC_ONLY'
        if cover<.98:verdict='INPUT_COVERAGE_BLOCKED'
        if n<max(12,cfg.required_folds*3):verdict='INSUFFICIENT_NONOVERLAP_CASES'
        evidence_pass=bool(mt['excess_cagr_pp']>=cfg.benchmark_excess_cagr_min_pp and
                           mt['ir'] is not None and mt['ir']>=cfg.ir_min and
                           mt_stress['excess_cagr_pp']>=cfg.benchmark_excess_cagr_min_pp and
                           mt['mdd']>=mt['benchmark_mdd']-cfg.max_extra_drawdown_pp/100 and cover>=.98 and
                           n>=max(12,cfg.required_folds*3) and ci[0] is not None and ci[0]>0)
        r={'status':verdict,'model':model,'horizon':h,'decision_count':len(data),
           'independent_cohorts':n,'data_coverage':cover,'blocked_exec':blocked,'blocked_fundamentals':fund_block,'blocked_other':other_block,
           'mean_stock_exposure':float(daily.alpha_weight.mean()),'trade_days_with_cost':int((daily.turnover_cost>0).sum()),
           'extra_roundtrip_stress_cagr_diff_pp':mt_stress['excess_cagr_pp'],
           'block_bootstrap_95_excess_log_annual_pp':ci,
           'candidate_performance_hurdle_met_retrospective':evidence_pass,
           'official_acceptance':'NOT_AUTHORIZED_CURRENT_SNAPSHOT_AND_EXPOSED_HISTORY',
           'final_tail_inspected':False,'new_independent_PIT_test':False,**mt}
        daily=daily.assign(model=model,horizon=h)
        return r,daily
    except Exception as exc:
        return {'status':'DAILY_PRICE_COVERAGE_BLOCKED','model':model,'horizon':h,
                'decision_count':len(data),'error':str(exc),'new_independent_PIT_test':False},pd.DataFrame()


def _chronological_folds(daily,decisions):
    """Exposed-development chronology blocks; not a pristine forward holdout.

    The model weights are frozen before the sweep; no fold outcome is fed back
    into fitting. Fold metrics are descriptive, not independent live evidence.
    """
    rows=[]
    if daily.empty or decisions.empty:return pd.DataFrame()
    for (model,h),sub in decisions.groupby(['model','horizon']):
        x=daily[(daily.model==model)&(daily.horizon==h)].sort_values('date')
        if x.empty:continue
        signals=sorted(pd.DatetimeIndex(sub.date.unique()))
        for i in range(0,len(signals),6):
            block=signals[i:i+6]
            if len(block)<3:continue
            begin=pd.Timestamp(block[0]);stop=pd.Timestamp(block[-1])
            if i+6<len(signals):stop=pd.Timestamp(signals[i+6])-pd.Timedelta(days=1)
            fragment=x[(x.date>=begin)&(x.date<=stop)].copy()
            if len(fragment)<15:continue
            m=_simple_fold_stat(fragment)
            rows.append({'model':model,'horizon':int(h),'fold':i//6+1,'start':begin.date().isoformat(),
                         'end':stop.date().isoformat(),'signals':len(block),**m,
                         'scope':'EXPOSED_DEVELOPMENT_CHRONOLOGICAL_BLOCK_NOT_BLIND_OOS'})
    return pd.DataFrame(rows)


def _simple_fold_stat(frame):
    a=float(frame.strategy_nav.iloc[-1]/frame.strategy_nav.iloc[0])
    b=float(frame.benchmark_nav.iloc[-1]/frame.benchmark_nav.iloc[0])
    n=max(len(frame)-1,1)
    ar=a**(252/n)-1;br=b**(252/n)-1
    return {'sessions':len(frame),'strategy_period_return':a-1,'0050_period_return':b-1,
            'excess_period_return_pp':100*(a-b),
            'annualized_excess_pp_diagnostic':100*(ar-br)}

def _save_excel(run,result,metrics,decisions,daily,qa,folds):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    wb=Workbook();s=wb.active;s.title='摘要'
    summary=[['Alpha Radar V23.1｜Benchmark-First 驗效報告','研究，不可直接下單'],
             ['engine',VERSION],['dataset_id',result['dataset_id']],['source_scope',result['source_scope']],
             ['final_tail_inspected',False],['independent_acceptance','NOT_YET'],['automatic_orders_authorized',False],
             ['cost_basis','Adjusted OHLC next-open; explicit stock/ETF fees and slippage'],
             ['warning','舊資料已反覆使用過，不能聲稱正式樣本外通過；資料來源及公告時間須另行稽核']]
    for row in summary:s.append(row)
    ms=wb.create_sheet('模型比較')
    keys=['model','horizon','status','decision_count','independent_cohorts','data_coverage','cagr','benchmark_cagr',
          'excess_cagr_pp','extra_roundtrip_stress_cagr_diff_pp','ir','mdd','candidate_performance_hurdle_met_retrospective','official_acceptance']
    ms.append(['模型','週期','狀態','決策月份','獨立持有期','資料覆蓋率','年化報酬','0050年化報酬','年化超額(百分點)',
               '成本壓力後超額(百分點)','IR','最大回撤','回顧績效門檻','正式驗收'])
    for i,m in enumerate(metrics,start=2):
        ms.append([m.get(k) for k in keys])
        ms.cell(i,9,f'=IF(OR(NOT(ISNUMBER(G{i})),NOT(ISNUMBER(H{i}))),"",(G{i}-H{i})*100)')
        for col in (7,8,12):ms.cell(i,col).number_format='0.0%'
        ms.cell(i,6).number_format='0.0%'
    def sheet_frame(name,frame,limit=100000):
        ws=wb.create_sheet(name)
        if frame is None or frame.empty:ws.append(['無資料']);return
        data=frame.copy().head(limit)
        for col in data.columns:
            if data[col].dtype=='object':
                data[col]=data[col].map(lambda x:json.dumps(x,ensure_ascii=False) if isinstance(x,(dict,list)) else str(x) if x is not None else '')
        ws.append(list(data.columns))
        for row in data.itertuples(index=False,name=None):
            ws.append([v.to_pydatetime().isoformat() if isinstance(v,pd.Timestamp) else
                       (None if pd.isna(v) else v) for v in row])
    sheet_frame('每月決策',decisions)
    sheet_frame('每日淨值',daily)
    sheet_frame('資料稽核',pd.DataFrame([qa]))
    sheet_frame('時序驗證區塊',folds)
    sheet_frame('研究設定',pd.DataFrame([asdict(V23Config(**result['settings']))]))
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines=False;ws.freeze_panes='B2' if ws.title=='模型比較' else 'A2'
        ws.row_dimensions[1].height=24
        for cell in ws[1]:
            cell.font=Font(name='Microsoft JhengHei',bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='183252');cell.alignment=Alignment(vertical='center')
        for j in range(1,min(ws.max_column,20)+1):
            ws.column_dimensions[get_column_letter(j)].width=min(38,max(14,len(str(ws.cell(1,j).value or ''))*1.4))
    wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True)
    out=run/'V23_1_VALIDATION.xlsx';wb.save(out)
    return out


def run_research(root,config=None,progress=lambda phase,message:None):
    root=Path(root).resolve();cfg=(config or V23Config()).validate()
    _,meta=load_dataset(root)
    release=load_release(root)
    models=release['models']
    expected={'version':VERSION,'dataset':meta['dataset_id'],'settings':asdict(cfg),
              'evidence_files':evidence_signature(root),'release':release,
              'code':{p:sha(Path(__file__).parent/p) for p in ('v231_runner.py','v231_engine.py','v231_factors.py','v231_data.py','v231_release.py')}}
    key=digest(expected)[:24]
    run=root/'v23'/'runs'/key;run.mkdir(parents=True,exist_ok=True)
    atomic_json(job_path(root),{'run':str(run),'run_id':key})
    if read_json(run/'state.json').get('status')=='COMPLETE' and (run/'V23_1_VALIDATION.zip').is_file():return run
    _stamp(run,status='RUNNING',phase='INITIALIZE',message='既有資料庫檢查中',pct=0)
    def status(phase,message):
        progress(phase,message)
        _stamp(run,status='RUNNING',phase=phase,message=message,pct={'PINNED_DATA':.07,'FEATURE_CACHE':.12,'LABEL_CACHE':.18,
                             'PIT_JOIN':.32,'SELECT':.47,'PORTFOLIO':.65,'REPORT':.91}.get(phase,.10))
    base=_read_base(root,cfg,status)
    dates,dev,holdout=_research_dates(base,cfg)
    events=load_events(root)
    signature=evidence_signature(root)
    status('PIT_JOIN','依公告日期建置PIT基本面因子；無資料保留缺漏')
    panel=_prepare_features(base,events,dev,cfg,status)
    status('SELECT','使用固定候選；不依已看過報酬自動調參')
    decisions,blocked=_decision_table(base,panel,dev,cfg,status,holdout['reserved_start'],models)
    if decisions.empty:raise DataError('沒有符合開發期且結束日已知的交易決策')
    # Stop labels if exits touch the holdout boundary, including 60-session exits.
    boundary=pd.Timestamp(holdout['reserved_start'])
    decisions=decisions[decisions.exit<boundary].copy()
    if decisions.empty:raise DataError('所有策略出場日進入保留期，不能評分')
    status('PORTFOLIO','用調整後價逐日重播資金與交易費用')
    prices=_subset_prices(base,decisions)
    metrics=[];curves=[]
    for h in (5,20,60):
        for m in models:
            status('PORTFOLIO',f'{h}日／{m}，只使用已暴露開發期')
            metric,curve=_validate_one(base,decisions,prices,m,h,cfg,status)
            metrics.append(metric)
            if not curve.empty:curves.append(curve)
    daily=pd.concat(curves,ignore_index=True) if curves else pd.DataFrame()
    folds=_chronological_folds(daily,decisions)
    source=meta.get('scope',{})
    source_scope=source.get('scope') if isinstance(source,dict) else str(source)
    qa={'dataset_id':meta['dataset_id'],'feature_cache_hit':base['cache_hit'],
        'feature_id':base['info'].get('feature_id'), 'evidence_signature':signature,
        'events_rows':{k:len(v) for k,v in events.items()},
        'selection_exec_label_blocks':blocked, 'historical_membership_verified':False,
        'price_revision_vintages_verified':False,'per_company_disclosure_independently_audited':False,
        'final_tail_inspected':False,'scope':source_scope,
        'methodology':'MONTHLY_DATE_T_SIGNAL_TPLUS1_OPEN_FIXED_HORZ_NONOVERLAP_ETF_0050_DAILY_NAV',
        'not_natively_weekly_5d_rebalancing':True}
    result={'version':VERSION,'research_only':True,'automatic_orders_authorized':False,
            'dataset_id':meta['dataset_id'],'source_scope':source_scope,'settings':asdict(cfg),
            'release_id':release['release_id'],'release_sha':release['release_sha'],
            'holdout_protocol':holdout,'development_dates':len(dev), 'evidence':signature,
            'models':metrics,'chronological_folds':len(folds),'data_quality':qa,'genuine_unseen_acceptance':False,
            'no_global_optimum_claimed':True,'report_generated_at':datetime.now(timezone.utc).isoformat()}
    status('REPORT','產出Excel、逐筆CSV、模型比較、資料品質及ZIP')
    atomic_json(run/'results.json',result)
    cv=decisions.copy()
    cv['selected']=cv.selected.map(lambda x:json.dumps(x,ensure_ascii=False))
    cv.to_csv(run/'monthly_decisions.csv',index=False,encoding='utf-8-sig')
    daily.to_csv(run/'daily_nav.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(metrics).to_csv(run/'candidate_comparison.csv',index=False,encoding='utf-8-sig')
    folds.to_csv(run/'chronological_folds.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame([qa]).to_csv(run/'data_quality.csv',index=False,encoding='utf-8-sig')
    atomic_json(run/'run_spec.json',expected)
    _save_excel(run,result,metrics,decisions,daily,qa,folds)
    short='\n'.join([f"- {x['horizon']}日 {x['model']}: {x['status']}, vs0050 CAGR差={x.get('excess_cagr_pp')}, IR={x.get('ir')}" for x in metrics])
    (run/'README_RESULTS.txt').write_text('V23.1｜研究報告，不可直接下單。\n'+
      '資料若是目前上市股票名冊及重建公告代理，不能稱為獨立PIT驗證。\n'+
      '最後24個月沒有評估；若先前其他模型已使用該時期，本次仍不得視為全新盲測。\n'+short,encoding='utf-8')
    names=['results.json','candidate_comparison.csv','chronological_folds.csv','monthly_decisions.csv','daily_nav.csv',
           'data_quality.csv','run_spec.json','V23_1_VALIDATION.xlsx','README_RESULTS.txt']
    zipout=run/'V23_1_VALIDATION.zip';tmp=zipout.with_suffix('.tmp')
    with zipfile.ZipFile(tmp,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for n in names:z.write(run/n,n)
    with zipfile.ZipFile(tmp) as z:
        bad=z.testzip()
        if bad:raise DataError('產出ZIP損壞：'+bad)
    tmp.replace(zipout)
    _stamp(run,status='COMPLETE',phase='DONE',message='已完成研究期驗證，正式獨立驗收仍未通過',pct=1.0)
    return run


def start(root,config=None):
    root=Path(root).resolve()
    with _ACTIVE_LOCK:
        prior=_ACTIVE.get(str(root))
        if prior and prior.is_alive():raise DataError('已有V23研究驗證工作在執行')
        def job():
            try:run_research(root,config)
            except Exception as exc:
                latest=read_json(job_path(root))
                if latest.get('run'):
                    _stamp(Path(latest['run']),status='ERROR',phase='ERROR',message=str(exc),
                          trace=traceback.format_exc()[-8000:])
        worker=threading.Thread(target=job,daemon=True,name='V231-research')
        _ACTIVE[str(root)]=worker;worker.start()
    return True
